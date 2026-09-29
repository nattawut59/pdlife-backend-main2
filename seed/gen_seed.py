#!/usr/bin/env python3
"""
Generates seed/seed.sql from docs/PDLIFE_Question_Bank_master_v1.xlsx.

Source of truth: the xlsx sheets, not this script's judgment. Every field is either a direct
column copy or a documented, narrow transformation (see comments below) — nothing here invents
data the workbook doesn't already contain. Where the workbook doesn't provide a value the schema
wants (e.g. checkin_templates.window_minutes), the seed leaves it NULL rather than guessing one.

Usage:
    python3 seed/gen_seed.py
Requires: openpyxl (pip install --user openpyxl)

The bootstrap admin is NOT generated here — see seed/gen_admin.mjs.

Output is idempotent — every INSERT is an upsert (ON CONFLICT DO UPDATE), so re-running this
after editing the workbook and re-generating is always safe to re-apply.
"""

import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(ROOT, "docs", "PDLIFE_Question_Bank_master_v1.xlsx")
OUT_PATH = os.path.join(ROOT, "seed", "seed.sql")

def esc(value) -> str:
    """SQL string literal, NULL-safe."""
    if value is None:
        return "NULL"
    s = str(value).strip()
    if s == "" or s == "-":
        return "NULL"
    return "'" + s.replace("'", "''") + "'"


def esc_bool(value) -> str:
    if value is None:
        return "NULL"
    return "TRUE" if value else "FALSE"


def assert_boolean_column_readable(wb, sheet: str, column: str) -> None:
    """หยุดทันทีถ้าคอลัมน์ boolean ทั้งคอลัมน์อ่านได้เป็นค่าว่าง

    เคยเกิดขึ้นจริงและไม่มีอะไรเตือนเลย: คอลัมน์ active เก็บเป็นสูตร =TRUE() ซึ่ง Excel เก็บ
    ทั้งตัวสูตรและผลลัพธ์ที่คำนวณไว้ (cached value) ส่วนสคริปต์นี้อ่านจากผลลัพธ์ (data_only=True)
    พอมีคนเปิดไฟล์ด้วย openpyxl แล้วบันทึกทับ openpyxl เขียนสูตรกลับได้แต่คำนวณผลลัพธ์ไม่เป็น
    จึงไม่เขียน cache กลับ — รอบต่อมาอ่านได้ None ทั้งคอลัมน์

    ผลคือ esc_bool(None) คืน "NULL" ซึ่งเขียนทับค่า DEFAULT true ของตาราง ทำให้ทุกคำถามและ
    ทุก template กลายเป็น inactive: สร้าง round ไม่ได้ ตอบคำถามไม่ได้ ทั้งระบบใช้งานไม่ได้

    เขียน NULL เงียบ ๆ แล้วปล่อยให้ไหลลงฐานข้อมูล แย่กว่าหยุดตรงนี้มาก
    """
    idx, rows = sheet_rows(wb, sheet)
    if column not in idx:
        return

    values = [row[idx[column]] for row in rows if any(cell is not None for cell in row)]
    if not values or any(v is not None for v in values):
        return

    sys.exit(
        "\n".join(
            [
                "",
                f"ERROR: ชีต '{sheet}' คอลัมน์ '{column}' อ่านได้ค่าว่างทั้ง {len(values)} แถว",
                "",
                "       สาเหตุที่เป็นไปได้มากที่สุด: คอลัมน์นี้เป็นสูตร แล้วค่าที่ Excel cache ไว้",
                "       หายไปตอนมีคนบันทึกไฟล์ด้วย openpyxl (openpyxl คำนวณสูตรไม่ได้)",
                "",
                "       วิธีแก้: เปิดไฟล์ใน Excel แล้วบันทึกใหม่",
                "               หรือเปลี่ยนสูตรในคอลัมน์นั้นเป็นค่าจริง (TRUE / FALSE)",
                "",
            ]
        )
    )


def esc_int(value) -> str:
    if value is None or str(value).strip() in ("", "-"):
        return "NULL"
    try:
        return str(int(value))
    except (TypeError, ValueError):
        return "NULL"


def esc_jsonb(raw) -> str:
    """Parses a JSON string from the sheet; falls back to '{}' if empty/unparseable
    (logging a warning rather than guessing at malformed data)."""
    if raw is None or str(raw).strip() in ("", "-"):
        return "'{}'::jsonb"
    try:
        parsed = json.loads(raw) if isinstance(raw, str) else raw
    except json.JSONDecodeError as e:
        print(f"WARNING: could not parse JSON ({e}): {raw!r} -> using {{}}", file=sys.stderr)
        return "'{}'::jsonb"
    compact = json.dumps(parsed, ensure_ascii=False)
    return "'" + compact.replace("'", "''") + "'::jsonb"


def sheet_rows(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header) if h is not None}
    return idx, rows[1:]


def build_options_from_response_options(wb) -> dict:
    """สร้าง options_json จากชีต response_options ซึ่งเป็นที่เดียวที่มี option_code จริง

    ชีต "Question Bank (dev-ready)" มีคอลัมน์ options_json อยู่ก็จริง แต่ข้างในมีแค่ value/label
    ที่เป็นข้อความไทย ไม่มี code — ขณะที่เงื่อนไข show_if ในชีต template_questions อ้างถึง
    รหัสภาษาอังกฤษ (yes, near_fall, …) ผลคือเงื่อนไขไม่มีทางเป็นจริงและคำถามปลายทางไม่เคยถูกถาม

    ชีต response_options มี option_code ครบทุกแถวอยู่แล้ว จึงใช้ชีตนั้นเป็นต้นทางแทน
    QID ที่ไม่มีในชีตนี้ (คำถามปลายเปิดที่ไม่มีตัวเลือก) จะถอยไปใช้คอลัมน์เดิม
    """
    if "response_options" not in wb.sheetnames:
        return {}

    idx, rows = sheet_rows(wb, "response_options")
    by_qid = {}
    for row in rows:
        qid = row[idx["QID"]]
        if not qid:
            continue
        code = row[idx["option_code"]]
        label = row[idx["option_label_th"]] or row[idx["option_value"]]
        score = row[idx["score"]]
        try:
            score = int(score) if score is not None and str(score).strip() != "" else None
        except (TypeError, ValueError):
            score = None
        by_qid.setdefault(str(qid).strip(), []).append(
            {"code": str(code).strip(), "label": str(label).strip(), "score": score}
        )
    return by_qid


def build_qid_to_code_map(wb) -> dict:
    """template_questions is the authoritative, complete mapping (schema §3: it's the single
    source of truth for what's actually asked). QID Mapping is used only to fill any gap."""
    mapping = {}

    idx, rows = sheet_rows(wb, "QID Mapping")
    for row in rows:
        qid = row[idx["QID (ของเรา)"]]
        code = row[idx["question_code (dev-ready)"]]
        if qid and code:
            mapping[qid] = code

    idx, rows = sheet_rows(wb, "template_questions")
    for row in rows:
        qid = row[idx["QID (ของเรา)"]]
        code = row[idx["question_code (dev-ready)"]]
        if qid and code:
            mapping[qid] = code  # template_questions wins on conflict

    return mapping


# question_bank.respondent — plain VARCHAR in the DB, but constrained at the app level to
# these three values (src/config/constants.ts RESPONDENT_POLICIES). Confirmed the sheet's
# respondent_lock only ever holds these three raw values before mapping.
RESPONDENT_LOCK_MAP = {
    "patient_or_caregiver": "both",
    "patient_only": "patient_only",
    "caregiver_preferred": "caregiver_preferred",
}


def gen_question_bank(wb, qid_to_code, options_by_qid) -> list:
    idx, rows = sheet_rows(wb, "Question Bank (dev-ready)")
    statements = []
    skipped = []

    for row in rows:
        qid = row[idx["QID"]]
        if not qid:
            continue
        code = qid_to_code.get(qid)
        if not code:
            skipped.append(qid)
            continue

        domain_code = qid.split("-")[0]
        respondent_lock = row[idx["respondent_lock"]]
        respondent = RESPONDENT_LOCK_MAP.get(respondent_lock, "both")

        cols = {
            "question_code": esc(code),
            "qid": esc(qid),
            "item_no": esc_int(row[idx["ข้อเดิม"]]),
            "domain_code": esc(domain_code),
            "domain_name_th": esc(row[idx["กลุ่ม"]]),
            "question_full_th": esc(row[idx["คำถาม"]]),
            "question_short_th": "NULL",  # not provided by the dev-ready sheet
            "answer_type": esc(row[idx["answer_type (dev)"]]),
            "options_json": (
                esc_jsonb(json.dumps(options_by_qid[str(qid).strip()], ensure_ascii=False))
                if str(qid).strip() in options_by_qid
                else esc_jsonb(row[idx["options_json"]])
            ),
            # question_bank.condition_json is a per-question *default* the template can override
            # (schema §3). The sheet's "เงื่อนไขการแสดง" column is free Thai text, not the
            # structured {show_if:{...}} format condition_json actually needs — parsing it would
            # be inventing structure the source doesn't provide, so this stays the DDL default.
            "condition_json": "'{}'::jsonb",
            "red_flag_json": esc_jsonb(row[idx["red_flag"]]),
            "respondent": esc(respondent),
            "ui_input_hint": esc(row[idx["ui_input_hint"]]),
            "required_level": esc(row[idx["required_level"]]),
            "summary_metric": esc(row[idx["summary_metric"]]),
            "mds_reference": esc(row[idx["MDS-UPDRS"]]),
            "license_source": esc(row[idx["license_source"]]),
            "terminology_binding": esc(row[idx["terminology_binding"]]),
            "mvp_phase": esc(row[idx["mvp_phase"]]),
            "active": esc_bool(row[idx["active"]]),
        }

        col_list = ", ".join(cols.keys())
        val_list = ", ".join(cols.values())
        update_list = ", ".join(f"{k} = EXCLUDED.{k}" for k in cols if k != "question_code")

        statements.append(
            f"INSERT INTO pdlife.question_bank ({col_list})\n"
            f"VALUES ({val_list})\n"
            f"ON CONFLICT (question_code) DO UPDATE SET {update_list}, updated_at = now();"
        )

    if skipped:
        print(f"WARNING: {len(skipped)} QIDs have no question_code mapping, skipped: {skipped}", file=sys.stderr)

    return statements


def gen_checkin_templates(wb) -> list:
    idx, rows = sheet_rows(wb, "checkin_templates")
    statements = []

    for row in rows:
        code = row[idx["template_code"]]
        if not code:
            continue

        cols = {
            "template_code": esc(code),
            "template_name_th": esc(row[idx["template_name_th"]]),
            "purpose": esc(row[idx["purpose"]]),
            "trigger_type": esc(row[idx["trigger_type"]]),
            # Not present as a clean number in the sheet ("recommended_time" is free text like
            # "หลังรับยา 30-60 นาที") — left NULL rather than parsed/guessed. The actual windows
            # are implemented directly in src/services/schedulerService.ts per schema §9.1's table.
            "window_minutes": "NULL",
            "max_questions_target": esc_int(row[idx["max_questions_target"]]),
            "notification_copy_th": esc(row[idx["notification_copy_th"]]),
            "mvp_phase": esc(row[idx["mvp_phase"]]),
            "active": esc_bool(row[idx["active"]]),
        }

        col_list = ", ".join(cols.keys())
        val_list = ", ".join(cols.values())
        update_list = ", ".join(f"{k} = EXCLUDED.{k}" for k in cols if k != "template_code")

        statements.append(
            f"INSERT INTO pdlife.checkin_templates ({col_list})\n"
            f"VALUES ({val_list})\n"
            f"ON CONFLICT (template_code) DO UPDATE SET {update_list};"
        )

    return statements


def gen_template_questions(wb) -> list:
    idx, rows = sheet_rows(wb, "template_questions")
    statements = []

    for row in rows:
        template_code = row[idx["template_code"]]
        question_code = row[idx["question_code (dev-ready)"]]
        if not template_code or not question_code:
            continue

        cols = {
            "template_code": esc(template_code),
            "question_code": esc(question_code),
            "sort_order": str(float(row[idx["sort_order"]])),
            "required": esc_bool(row[idx["required"]]),
            "condition_json": esc_jsonb(row[idx["condition_json"]]),
            "ui_note_th": esc(row[idx["ui_note_th"]]),
        }

        col_list = ", ".join(cols.keys())
        val_list = ", ".join(cols.values())
        update_list = ", ".join(
            f"{k} = EXCLUDED.{k}" for k in ("sort_order", "required", "condition_json", "ui_note_th")
        )

        statements.append(
            f"INSERT INTO pdlife.template_questions ({col_list})\n"
            f"VALUES ({val_list})\n"
            f"ON CONFLICT (template_code, question_code) DO UPDATE SET {update_list};"
        )

    return statements


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    # ตรวจก่อนสร้างอะไรทั้งนั้น — คอลัมน์พวกนี้ควบคุมว่าคำถาม/template ใช้งานได้หรือไม่
    assert_boolean_column_readable(wb, "Question Bank (dev-ready)", "active")
    assert_boolean_column_readable(wb, "checkin_templates", "active")
    assert_boolean_column_readable(wb, "template_questions", "required")

    qid_to_code = build_qid_to_code_map(wb)
    options_by_qid = build_options_from_response_options(wb)

    parts = []
    parts.append("-- Generated by seed/gen_seed.py from docs/PDLIFE_Question_Bank_master_v1.xlsx")
    parts.append("-- Do not hand-edit — re-run the generator instead. Safe to re-apply (upserts).")
    parts.append("BEGIN;")

    parts.append("\n-- ========== bootstrap admin ==========")
    parts.append(
        "-- Deliberately not here. This file is committed, so any password in it is known to\n"
        "-- everyone with repo access. Generate the first admin separately:\n"
        "--\n"
        "--     node seed/gen_admin.mjs        (prints the password once, writes admin.generated.sql)\n"
        "--\n"
        "-- then apply seed/admin.generated.sql after this file."
    )

    parts.append("\n-- ========== question_bank ==========")
    parts.extend(gen_question_bank(wb, qid_to_code, options_by_qid))

    parts.append("\n-- ========== checkin_templates ==========")
    parts.extend(gen_checkin_templates(wb))

    parts.append("\n-- ========== template_questions ==========")
    parts.extend(gen_template_questions(wb))

    parts.append("\nCOMMIT;")

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write("\n\n".join(parts) + "\n")

    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
